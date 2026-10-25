import copy
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

# Insere na Aula 5, antes do exemplo da Viagem de Finalistas, o slide
# "O Modelo de Atividade em Pormenor" com a figura de atividade_fuks.py.
# Copia o cabeçalho e a linha de texto do slide "Exemplo: o Modelo
# Preenchido". Depois correr numerar_slides.py. Só com o deck fechado.
DECK = sys.argv[1]
IMG = ("/Users/jpavao/QProjects/Sebenta-PSC/assets/images/cap04/slides/"
       "teoria-atividade-fuks-slide.png")
TITULO = "O Modelo de Atividade em Pormenor"
CREDITO = "Adaptado de Fuks et al. (2011), a partir de Engeström (1987)."
NOTAS = (
    "• Este esquema mostra o mesmo modelo do slide anterior, agora com exemplos de cada elemento\n\n"
    "• Ler cada canto com a turma: quem pode ser o sujeito, que artefactos usa, que regras, que comunidade, que divisão de trabalho\n\n"
    "• Os quatro triângulos interiores são as quatro partes de qualquer atividade em sociedade: "
    "produção (fazer, com artefactos), troca (comunicar, sob regras), distribuição (repartir, pela divisão de trabalho) e consumo (usar o que se produziu)\n\n"
    "• Ligação aos 3C, que aparecem mais à frente: produção = cooperação, troca = comunicação, distribuição = coordenação\n\n"
    "• A seguir, aplicar o modelo a uma situação concreta: a viagem de finalistas"
)

p = Presentation(DECK)
if any(sh.has_text_frame and sh.text_frame.text == TITULO for s in p.slides for sh in s.shapes):
    sys.exit("o slide já existe")
modelo = next(s for s in p.slides
              if any(sh.has_text_frame and sh.text_frame.text == "Exemplo: o Modelo Preenchido"
                     for sh in s.shapes))
viagem = next(i for i, s in enumerate(p.slides)
              if any(sh.has_text_frame and sh.text_frame.text == "Exemplo: a Viagem de Finalistas"
                     for sh in s.shapes))

novo = p.slides.add_slide(modelo.slide_layout)
for ph in list(novo.placeholders):
    ph._element.getparent().remove(ph._element)
for sh in modelo.shapes:
    if sh.shape_type == 13 or sh.name == "NumeroSlide":
        continue
    novo.shapes._spTree.append(copy.deepcopy(sh._element))

for sh in novo.shapes:
    if not sh.has_text_frame:
        continue
    if sh.text_frame.text == "Exemplo: o Modelo Preenchido":
        sh.text_frame.paragraphs[0].runs[0].text = TITULO
    elif sh.text_frame.text.startswith("Ação reversa"):
        r = sh.text_frame.paragraphs[0].runs[0]
        r.text = CREDITO
        r.font.italic = True

x, y, w, h = 661475, 1450000, 10868745, 4400000
iw, ih = Image.open(IMG).size
k = min(w / iw, h / ih)
nw, nh = int(iw * k), int(ih * k)
novo.shapes.add_picture(IMG, Emu(x + (w - nw) // 2), Emu(y + (h - nh) // 2), Emu(nw), Emu(nh))

ns = novo.notes_slide
if ns.notes_placeholder is None:
    ns.shapes._spTree.append(copy.deepcopy(modelo.notes_slide.notes_placeholder._element))
ns.notes_text_frame.text = NOTAS

lst = p.slides._sldIdLst
el = list(lst)[-1]
lst.remove(el)
lst.insert(viagem, el)
p.save(DECK)
print("inserido como slide", viagem + 1, "de", len(p.slides))

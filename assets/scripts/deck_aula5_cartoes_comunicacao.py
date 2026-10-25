import copy
import sys
from pptx import Presentation
from pptx.util import Emu

# Aula 5: o slide "Comunicação: A Mensagem como Elemento Central" passa a
# ter os três atributos em cartões, como os slides de Formação de Grupos e
# de Coordenação. O cartão é a mesma imagem usada no slide de Coordenação.
# Só com o deck fechado.
DECK = sys.argv[1]

p = Presentation(DECK)


def slide_com(titulo):
    return next(s for s in p.slides
                if any(sh.has_text_frame and sh.text_frame.text == titulo for sh in s.shapes))


com = slide_com("Comunicação: A Mensagem como Elemento Central")
coord = slide_com("Coordenação: O Plano de Trabalho como Elemento Central")
if any(sh.shape_type == 13 for sh in com.shapes):
    sys.exit("o slide já tem cartões")

# o cartão é uma imagem só SVG (svgBlip): copia-se o elemento e liga-se o
# slide à mesma parte de imagem
cartao = next(sh for sh in coord.shapes if sh.shape_type == 13)
SVG = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
svg_rid = cartao._element.find(".//" + SVG).get(R_EMBED)
rid = com.part.relate_to(coord.part.related_part(svg_rid),
                         coord.part.rels[svg_rid].reltype)
sh = {s.name: s for s in com.shapes}

topo, passo = 2650000, 1072158
dy_titulo, dy_corpo = 168870, 538857
for i, (tit, corpo) in enumerate([("Text 4", "Text 5"), ("Text 6", "Text 7"), ("Text 8", "Text 9")]):
    y = topo + i * passo
    el = copy.deepcopy(cartao._element)
    el.find(".//" + SVG).set(R_EMBED, rid)
    el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}off").set("y", str(y))
    el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("id", str(900 + i))
    # o cartão fica por trás do texto
    sh[tit]._element.addprevious(el)
    for nome, dy in [(tit, dy_titulo), (corpo, dy_corpo)]:
        s = sh[nome]
        s.left, s.top, s.width = Emu(906539), Emu(y + dy), Emu(5399250)

sh["Text 10"].top = Emu(topo + 2 * passo + cartao.height + 110000)
p.save(DECK)
print("ok", list(p.slides).index(com) + 1)

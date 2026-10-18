import copy
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

# Aula 5: a seguir a cada slide de texto da ontologia, um slide com o
# diagrama correspondente a toda a largura (slide_cap04_ontologias.py);
# a ilustração Gamma do slide de texto é retirada.
DECK = sys.argv[1]
IMGDIR = "/Users/jpavao/QProjects/Sebenta-PSC/assets/images/cap04/slides"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

PARTS = [
    ("Formação de Grupos: O Participante", "grupos", "Ontologia sobre Formação de Grupos",
     "o Participante", "o Participante assume Tarefas; a Confiança é atributo do Participante"),
    ("Comunicação: A Mensagem", "comunicacao", "Ontologia sobre Comunicação",
     "a Mensagem", "o Emissor envia a Mensagem; o Meio de Transmissão é atributo da Mensagem"),
    ("Coordenação: O Plano de Trabalho", "coordenacao", "Ontologia sobre Coordenação",
     "o Plano de Trabalho", "o Plano de Trabalho define Prazos; o Acompanhamento controla o Estado"),
    ("Cooperação: A Atividade", "cooperacao", "Ontologia sobre Cooperação",
     "a Atividade", "a Atividade gera um Produto; a Tarefa é parte da Atividade"),
]

p = Presentation(DECK)
lst = p.slides._sldIdLst


def title_of(slide):
    for sh in slide.shapes:
        if sh.name == "Text 2" and sh.has_text_frame:
            return sh.text_frame.text
    return ""


def set_text(sh, text):
    paras = sh.text_frame._txBody.findall(A + "p")
    for extra in paras[1:]:
        extra.getparent().remove(extra)
    runs = paras[0].findall(A + "r")
    for r in runs[1:]:
        paras[0].remove(r)
    runs[0].find(A + "t").text = text


for prefix, name, title, centre, examples in PARTS:
    slides = list(p.slides)
    ref = next(s for s in slides if title_of(s).startswith(prefix))
    pos = slides.index(ref)

    # retirar a ilustração Gamma (a imagem grande e quadrada)
    # quando a ilustração estava à esquerda, o texto passa para o lugar dela
    for sh in list(ref.shapes):
        if sh.shape_type == 13 and sh.width > 4000000 and sh.height > 4000000:
            pic_left = sh.left
            sh._element.getparent().remove(sh._element)
            if pic_left < 1000000:
                dx = 5723290 - pic_left
                for other in ref.shapes:
                    if other.left >= 5700000:
                        other.left = Emu(other.left - dx)

    s = p.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    shapes = {sh.name: sh for sh in ref.shapes}
    for nm in ("Shape 0", "Text 1", "Text 2"):
        s.shapes._spTree.append(copy.deepcopy(shapes[nm]._element))
    set_text(s.shapes[-1], title)

    img = f"{IMGDIR}/ontologia-{name}-slide.png"
    iw, ih = Image.open(img).size
    top, max_w, max_h = 1350000, 10868745, 5250000
    w = max_w
    h = int(w * ih / iw)
    if h > max_h:
        h = max_h
        w = int(h * iw / ih)
    s.shapes.add_picture(img, Emu((p.slide_width - w) // 2), Emu(top), Emu(w), Emu(h))

    ns = s.notes_slide
    if ns.notes_placeholder is None:
        ns.shapes._spTree.append(copy.deepcopy(ref.notes_slide.notes_placeholder._element))
    ns.notes_text_frame.text = (
        f"• Ler o diagrama a partir do elemento central, a cheio: {centre}\n\n"
        f"• Cada seta lê-se como uma frase, por exemplo: {examples}\n\n"
        "• As elipses a tracejado são conceitos que aparecem também noutra parte da ontologia\n\n"
        "• As caixas de nota (canto dobrado) dão exemplos dos valores possíveis")

    new = list(lst)[-1]
    lst.remove(new)
    lst.insert(pos + 1, new)

p.save(DECK)
print("slides:", len(p.slides))
for i, sl in enumerate(p.slides, 1):
    t = title_of(sl)
    if "Ontologia" in t or any(t.startswith(x[0]) for x in PARTS) or "Tuckman" in t:
        print(i, t)

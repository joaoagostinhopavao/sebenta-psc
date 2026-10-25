import copy
import sys
from pptx import Presentation
from pptx.util import Emu, Pt

# Aula 5: o slide "Aplicação à Sala de Aula" (redundante depois da Viagem
# de Finalistas) passa a ser a ponte "Da Atividade à Colaboração"; o
# Modelo 3C só é nomeado no fim, como passagem para o slide seguinte. Reaproveita o
# cabeçalho, três ícones e os estilos de texto do próprio slide.
# Só com o deck fechado.
DECK = sys.argv[1]

LINHAS = [  # (ícone, título, corpo)
    ("Image 3", "Troca  =  Comunicação",
     "As discussões e as sondagens no grupo de WhatsApp: a turma negoceia e assume "
     "compromissos (máximo de 350 € por pessoa)."),
    ("Image 4", "Distribuição  =  Coordenação",
     "A divisão de trabalho (Ana trata do alojamento, Bruno do transporte…) e os "
     "prazos (quota paga até 15 de março)."),
    ("Image 2", "Produção  =  Cooperação",
     "O trabalho feito em conjunto no espaço partilhado: a folha de cálculo dos "
     "pagamentos e as reservas."),
]
NOTAS = (
    "• Ponte entre os dois modelos: três das quatro partes da atividade em sociedade (slide 9) "
    "são as três dimensões do Modelo 3C: produção = cooperação, distribuição = coordenação, "
    "troca = comunicação (Fuks et al., 2011)\n\n"
    "• A Teoria da Atividade descreve a estrutura da atividade (quem, sobre quê, com que "
    "ferramentas, com que regras); o Modelo 3C descreve a dinâmica da colaboração\n\n"
    "• Ler a viagem pelos 3C com a turma, linha a linha\n\n"
    "• Terminar com o imprevisto do hotel: a cooperação exige nova comunicação, que exige "
    "coordenação\n\n"
    "• Só no fim se dá nome ao Modelo 3C, que o slide seguinte apresenta em pormenor"
)

p = Presentation(DECK)
slide = next(s for s in p.slides
             if any(sh.has_text_frame and sh.text_frame.text == "Aplicação à Sala de Aula"
                    for sh in s.shapes))
sh = {s.name: s for s in slide.shapes}

sh["Text 2"].text_frame.paragraphs[0].runs[0].text = "Da Atividade à Colaboração"
sh["Text 3"].text_frame.paragraphs[0].runs[0].text = (
    "Três das quatro partes da atividade correspondem a três dimensões da colaboração. "
    "Na viagem de finalistas:")
sh["Text 3"].width = Emu(10868745)

titulos, corpos = ["Text 4", "Text 6", "Text 8"], ["Text 5", "Text 7", "Text 9"]
topo, passo = 2713732, 850000
for i, (icone, titulo, corpo) in enumerate(LINHAS):
    y = topo + i * passo
    sh[icone].left, sh[icone].top = Emu(661475), Emu(y)
    t, c = sh[titulos[i]], sh[corpos[i]]
    t.left, t.top, t.width = Emu(1198830), Emu(y + 56852), Emu(10331390)
    c.left, c.top, c.width, c.height = Emu(1198830), Emu(y + 414536), Emu(10331390), Emu(264616)
    t.text_frame.paragraphs[0].runs[0].text = titulo
    c.text_frame.paragraphs[0].runs[0].text = corpo

for nome in ["Image 0", "Image 1", "Image 5", "Text 10", "Text 11", "Text 12", "Text 13",
             "Text 14", "Text 15"]:
    sh[nome]._element.getparent().remove(sh[nome]._element)

# remate: o imprevisto que leva ao ciclo do slide seguinte
el = copy.deepcopy(sh["Text 3"]._element)
slide.shapes._spTree.append(el)
fim = slide.shapes[-1]
fim.top, fim.height = Emu(topo + 3 * passo + 80000), Emu(900000)
par = fim.text_frame.paragraphs[0]
r1 = par.runs[0]
r1.text = "E se o hotel cancelar? "
r1.font.bold = True
r2 = copy.deepcopy(r1._r)
r1._r.addnext(r2)
par.runs[1].text = ("A comissão volta ao WhatsApp para decidir (comunicação) e a Ana "
                    "reorganiza o alojamento (coordenação): as três dimensões alimentam-se "
                    "umas às outras.")
par.runs[1].font.bold = False
# segundo parágrafo: só aqui se nomeia o modelo
p2 = copy.deepcopy(par._p)
par._p.addnext(p2)
par2 = fim.text_frame.paragraphs[1]
for r, (texto, negrito) in zip(par2.runs, [("Estas três dimensões estão na base de um modelo "
                                            "de colaboração, o ", False)]):
    r.text, r.font.bold = texto, negrito
r3 = copy.deepcopy(par2.runs[1]._r)
par2.runs[1]._r.addnext(r3)
par2.runs[0].text, par2.runs[0].font.bold = ("Estas três dimensões estão na base de um "
                                             "modelo de colaboração, o "), False
par2.runs[1].text, par2.runs[1].font.bold = "Modelo 3C", True
par2.runs[2].text, par2.runs[2].font.bold = ".", False
par2.space_before = Pt(8)
fim.text_frame.word_wrap = True

slide.notes_slide.notes_text_frame.text = NOTAS
p.save(DECK)
print("ok", p.slides.index(slide) + 1)

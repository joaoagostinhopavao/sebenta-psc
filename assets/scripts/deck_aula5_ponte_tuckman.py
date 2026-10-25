import copy
import sys
from pptx import Presentation

# Aula 5: slide-ponte entre "3C vs. Comando e Controlo" e o Modelo de
# Tuckman. Conta a vida da comissão da viagem mês a mês, com o C que pesa
# mais em cada fase, e só no fim nomeia o modelo. Copia o formato do slide
# "Exemplo: a Viagem de Finalistas". Depois correr numerar_slides.py.
# Só com o deck fechado.
DECK = sys.argv[1]
TITULO = "Um Grupo Colabora Bem desde o Início?"
ETIQUETA = "COLABORAÇÃO"
PARAGRAFOS = [
    "No Comando e Controlo, a organização chega pronta. Num grupo que colabora, é o próprio "
    "grupo que a tem de construir, e isso leva tempo. Na comissão da viagem:",
    "1. Setembro: os quatro mal se conhecem e falam sobretudo do que querem para a viagem "
    "(comunicação).",
    "2. Outubro: discutem quem faz o quê; há atritos até cada um ficar com a sua parte "
    "(comunicação).",
    "3. Novembro: combinam como vão trabalhar: votar por sondagem e pagar a quota até 15 de "
    "março (coordenação).",
    "4. De janeiro a abril: cada um trata da sua parte, quase sem precisar de reuniões "
    "(cooperação).",
    "5. Depois da viagem: a comissão desfaz-se.",
    "O grupo passou por etapas e, em cada uma, pesou mais um dos três C. Bruce Tuckman "
    "descreveu estas etapas: é o Modelo de Tuckman.",
]
NOTAS = (
    "• Fio condutor entre o 3C e o Tuckman: no C2 a organização é imposta desde o início; num "
    "grupo colaborativo é construída pelo próprio grupo, ao longo do tempo\n\n"
    "• Contar a vida da comissão da viagem mês a mês e perguntar à turma, em cada passo, qual "
    "dos três C está a pesar mais\n\n"
    "• Os três C estão sempre presentes; o que muda é qual exige mais atenção\n\n"
    "• Só no fim se nomeia o Modelo de Tuckman, que o slide seguinte apresenta com o nome de "
    "cada etapa"
)

p = Presentation(DECK)
if any(sh.has_text_frame and sh.text_frame.text == TITULO for s in p.slides for sh in s.shapes):
    sys.exit("o slide já existe")
slides = list(p.slides)
modelo = next(s for s in slides if any(sh.has_text_frame and sh.text_frame.text ==
                                       "Exemplo: a Viagem de Finalistas" for sh in s.shapes))
c2 = next(i for i, s in enumerate(slides) if any(sh.has_text_frame and sh.text_frame.text ==
                                                 "3C vs. Comando e Controlo" for sh in s.shapes))

novo = p.slides.add_slide(modelo.slide_layout)
for ph in list(novo.placeholders):
    ph._element.getparent().remove(ph._element)
for sh in modelo.shapes:
    if sh.name != "NumeroSlide":
        novo.shapes._spTree.append(copy.deepcopy(sh._element))

sh = {s.name: s for s in novo.shapes}
sh["Text 1"].text_frame.paragraphs[0].runs[0].text = ETIQUETA
sh["Text 2"].text_frame.paragraphs[0].runs[0].text = TITULO
paras = sh["Text 3"].text_frame.paragraphs
assert len(paras) == len(PARAGRAFOS)
for para, texto in zip(paras, PARAGRAFOS):
    para.runs[0].text = texto
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)

ns = novo.notes_slide
if ns.notes_placeholder is None:
    ns.shapes._spTree.append(copy.deepcopy(modelo.notes_slide.notes_placeholder._element))
ns.notes_text_frame.text = NOTAS

lst = p.slides._sldIdLst
el = list(lst)[-1]
lst.remove(el)
lst.insert(c2 + 1, el)
p.save(DECK)
print("inserido como slide", c2 + 2, "de", len(p.slides))

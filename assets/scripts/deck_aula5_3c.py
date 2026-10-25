import copy
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

# Slide "O Modelo 3C de Colaboração" da Aula 5: troca os cartões de texto
# e a imagem pela figura de modelo_3c_fuks.py, com o crédito por baixo
# (copiado do slide "O Modelo de Atividade em Pormenor"). Mantém o
# cabeçalho e o número. Só com o deck fechado.
DECK = sys.argv[1]
IMG = "/Users/jpavao/QProjects/Sebenta-PSC/assets/images/cap04/slides/modelo-3c-fuks-slide.png"
TITULO = "O Modelo 3C de Colaboração"
NOTAS_EXTRA = (
    "\n\n• Ler o esquema como um ciclo: ao comunicar, as pessoas negoceiam e assumem compromissos, "
    "que são geridos pela coordenação; a coordenação organiza as tarefas para a cooperação; "
    "e os imprevistos da cooperação exigem nova comunicação\n\n"
    "• Ao centro, a perceção: é através dela que cada um tem retorno das suas próprias ações "
    "e conhece as ações dos colegas, nas três dimensões\n\n"
    "• A etimologia ajuda a memorizar: comum-ação (tornar comum), co-ordem-ação (ordenar em conjunto), "
    "co-operar-ação (operar em conjunto)\n\n"
    "• Perceção (awareness) = ter noção do que os outros estão a fazer: cada um sabe o suficiente "
    "do que os outros fazem para saber o que ele próprio deve fazer. Não é saber tudo: na viagem, "
    "a Carla não precisa dos pormenores do autocarro que o Bruno está a negociar, mas precisa de saber "
    "que o transporte está a ser tratado, quanto vai custar e quem já pagou\n\n"
    "• As setas nos dois sentidos: o que acontece em cada C alimenta a perceção (uma mensagem, uma tarefa "
    "atribuída, uma alteração na folha de cálculo) e a perceção orienta cada C"
)

p = Presentation(DECK)
slide = next(s for s in p.slides
             if any(sh.has_text_frame and sh.text_frame.text == TITULO for sh in s.shapes))
fonte = next(sh for s in p.slides for sh in s.shapes
             if sh.has_text_frame and sh.text_frame.text.startswith("Adaptado de Fuks"))

manter = {"Shape 0", "Text 1", "Text 2", "NumeroSlide"}
for sh in list(slide.shapes):
    if sh.name not in manter:
        sh._element.getparent().remove(sh._element)

x, y, w, h = 661475, 1400000, 10868745, 4450000
iw, ih = Image.open(IMG).size
k = min(w / iw, h / ih)
nw, nh = int(iw * k), int(ih * k)
slide.shapes.add_picture(IMG, Emu(x + (w - nw) // 2), Emu(y + (h - nh) // 2), Emu(nw), Emu(nh))

el = copy.deepcopy(fonte._element)
slide.shapes._spTree.append(el)
cred = slide.shapes[-1]
cred.text_frame.paragraphs[0].runs[0].text = "Adaptado de Fuks et al. (2011)."

tf = slide.notes_slide.notes_text_frame
if "Ler o esquema como um ciclo" in tf.text:
    tf.text = tf.text[:tf.text.index("\n\n• Ler o esquema como um ciclo")]
tf.text = tf.text.rstrip() + NOTAS_EXTRA
p.save(DECK)
print("ok", p.slides.index(slide) + 1)

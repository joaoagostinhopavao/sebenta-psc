import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

# Troca as imagens do Modelo de Atividade nos slides 8 ("Como se Realiza
# uma Atividade Mediada?") e 10 ("Exemplo: o Modelo Preenchido") da Aula 5
# pelas versões de Engeström (slide_cap04_atividade*.py). Cada imagem é
# ajustada à caixa indicada, mantendo a proporção. Só com o deck fechado.
DECK = sys.argv[1]
DIR = "/Users/jpavao/QProjects/Sebenta-PSC/assets/images/cap04/slides/"
TROCAS = [
    ("Como se Realiza uma Atividade Mediada?", DIR + "teoria-atividade-slide.png",
     None),
    ("Exemplo: o Modelo Preenchido", DIR + "teoria-atividade-exemplo-slide.png",
     (661475, 1450000, 10868745, 4400000)),
]

p = Presentation(DECK)
for titulo, img, caixa in TROCAS:
    slide = next(s for s in p.slides
                 if any(sh.has_text_frame and sh.text_frame.text == titulo
                        for sh in s.shapes))
    old = next(sh for sh in slide.shapes if sh.shape_type == 13)
    x, y, w, h = caixa or (old.left, old.top, old.width, old.height)
    old._element.getparent().remove(old._element)
    iw, ih = Image.open(img).size
    escala = min(w / iw, h / ih)
    nw, nh = int(iw * escala), int(ih * escala)
    slide.shapes.add_picture(img, Emu(x + (w - nw) // 2), Emu(y + (h - nh) // 2),
                             Emu(nw), Emu(nh))
    print("ok", p.slides.index(slide) + 1)
p.save(DECK)

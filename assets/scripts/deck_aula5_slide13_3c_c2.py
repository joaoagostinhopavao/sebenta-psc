import sys
from pptx import Presentation
from pptx.util import Emu

# Troca a imagem do slide "3C vs. Comando e Controlo" (Aula 5) pela
# gerada em slide_cap04_3c_c2.py, ocupando toda a largura útil do slide.
DECK = sys.argv[1]
IMG = "/Users/jpavao/QProjects/Sebenta-PSC/assets/images/cap04/slides/3c-vs-c2-slide.png"

p = Presentation(DECK)
slide = next(s for s in p.slides
             if any(sh.has_text_frame and sh.text_frame.text == "3C vs. Comando e Controlo"
                    for sh in s.shapes))
old = next(sh for sh in slide.shapes if sh.shape_type == 13)
top = old.top
old._element.getparent().remove(old._element)

w = Emu(10868745)
h = Emu(int(w * 1196 / 2666))
pic = slide.shapes.add_picture(IMG, Emu(661475), top, w, h)
p.save(DECK)
print("ok", p.slides.index(slide) + 1)

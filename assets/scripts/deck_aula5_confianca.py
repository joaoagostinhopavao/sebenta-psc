import sys
from pptx import Presentation

# Aula 5, slide "Formação de Grupos": confiança passa a ser definida como
# disposição para depender dos outros (Mayer et al., 1995), distinta das
# crenças, em linha com o Cap04.
DECK = sys.argv[1]

CARD_OLD = ("A crença partilhada de que se pode contar com os outros participantes "
            "para cumprir os compromissos assumidos.")
CARD_NEW = ("A disposição para depender dos outros, com base na expectativa de que "
            "cumprirão os compromissos assumidos.")
NOTE_OLD = "Confiança: a crença partilhada de que os outros vão cumprir os compromissos assumidos"
NOTE_NEW = ("Confiança: a disposição para depender dos outros, aceitando ficar exposto ao que "
            "fizerem, com base na expectativa de que cumprirão os compromissos assumidos "
            "(Mayer, Davis e Schoorman, 1995)\n\n"
            "Não confundir com as crenças, outro atributo do participante no diagrama: as crenças "
            "são o que o participante pensa (por exemplo, que um colega é competente); a confiança "
            "é a disposição para depender dele e assenta nessas crenças. Pode haver crença sem "
            "confiança: acreditar que o colega é competente e, mesmo assim, não lhe entregar parte "
            "do trabalho")

p = Presentation(DECK)
done = 0
for s in p.slides:
    frames = [sh.text_frame for sh in s.shapes if sh.has_text_frame]
    if s.has_notes_slide:
        frames.append(s.notes_slide.notes_text_frame)
    for tf in frames:
        for para in tf.paragraphs:
            for r in para.runs:
                if CARD_OLD in r.text:
                    r.text = r.text.replace(CARD_OLD, CARD_NEW)
                    done += 1
                if NOTE_OLD in r.text:
                    r.text = r.text.replace(NOTE_OLD, NOTE_NEW)
                    done += 1
assert done == 2, done
p.save(DECK)
print("ok")

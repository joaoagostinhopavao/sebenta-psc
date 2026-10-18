import copy
import sys
from pptx import Presentation
from pptx.util import Emu

# Aula 4b, slide do caminho/distância/diâmetro: o diâmetro é a maior das
# distâncias (cada uma o caminho MAIS CURTO), com o exemplo I-G = 5 e não 6.
DECK = sys.argv[1]
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

DEF_OLD = (" — a maior distância entre quaisquer dois nós da rede; indica quanto tempo, "
           "no pior caso, a informação demora a atravessá-la")
DEF_NEW = (" — a maior das distâncias da rede, sendo cada distância o caminho mais curto "
           "entre dois nós; indica quanto tempo, no pior caso, a informação demora a atravessá-la")
NOTE_OLD = "• Diâmetro: a maior distância da rede. Na turma é 5, de I a H (I, A, C, E, F, H) e também de I a G"
NOTE_NEW = ("• Diâmetro: a maior das distâncias da rede, e cada distância é o caminho MAIS CURTO entre "
            "dois nós. Não é o caminho mais comprido que se consegue traçar: dando voltas, há quase "
            "sempre um caminho mais longo. Na turma é 5, de I a H (I, A, C, E, F, H) e também de I a G "
            "(I, A, C, E, F, G). O percurso I, A, C, E, F, H, G tem 6 passos mas não conta, porque F e G "
            "estão ligados diretamente")

p = Presentation(DECK)
slide = next(s for s in p.slides
             if any(sh.has_text_frame and sh.text_frame.text.startswith("Diâmetro — ")
                    for sh in s.shapes))
done = 0
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            if r.text == DEF_OLD:
                r.text = DEF_NEW
                sh.height = Emu(1150000)
                done += 1
    if sh.text_frame.text.startswith("No exemplo da rede da turma"):
        # segundo parágrafo com o contraexemplo I-G
        txBody = sh.text_frame._txBody
        p0 = txBody.findall(A + "p")[0]
        p1 = copy.deepcopy(p0)
        runs = p1.findall(A + "r")
        for r in runs[2:]:
            p1.remove(r)
        runs[0].find(A + "t").text = "De I a G também são 5 passos, e não 6: "
        p1.remove(runs[1])
        runs[0].find(A + "t").text = ("De I a G também são 5 passos, e não 6: I → A → C → E → F → H → G "
                                      "é um caminho, mas não o mais curto, porque F e G estão "
                                      "ligados diretamente.")
        p0.addnext(p1)
        sh.top = Emu(5520000)
        sh.width = Emu(10868745)
        sh.height = Emu(1000000)
        done += 1
notes = slide.notes_slide.notes_text_frame
for para in notes.paragraphs:
    for r in para.runs:
        if NOTE_OLD in r.text:
            r.text = r.text.replace(NOTE_OLD, NOTE_NEW)
            done += 1
assert done == 3, done
p.save(DECK)
print("ok")

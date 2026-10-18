import copy
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import MSO_AUTO_SIZE

# Acrescenta à Aula 5, a seguir ao slide "Formação de Grupos", um slide com
# o Modelo de Tuckman (imagem de slide_cap04_tuckman.py).
DECK = sys.argv[1]
IMG = "/Users/jpavao/QProjects/Sebenta-PSC/assets/images/cap04/slides/tuckman-slide.png"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

p = Presentation(DECK)
idx = next(i for i, s in enumerate(p.slides)
           if any(sh.has_text_frame and sh.text_frame.text.startswith("Formação de Grupos: O Participante")
                  for sh in s.shapes))
ref = p.slides[idx]
cap_src = next(sh for s in p.slides for sh in s.shapes
               if sh.has_text_frame and sh.text_frame.text.startswith("Só há "))
shapes = {sh.name: sh for sh in ref.shapes}

s = p.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)


def clone(sh):
    s.shapes._spTree.append(copy.deepcopy(sh._element))
    return s.shapes[-1]


def set_text(sh, text, size=None):
    paras = sh.text_frame._txBody.findall(A + "p")
    for extra in paras[1:]:
        extra.getparent().remove(extra)
    runs = paras[0].findall(A + "r")
    for r in runs[1:]:
        paras[0].remove(r)
    runs[0].find(A + "t").text = text
    rpr = runs[0].find(A + "rPr")
    rpr.set("b", "0")
    if size:
        rpr.set("sz", str(size))


pill = clone(shapes["Shape 0"])
tag = clone(shapes["Text 1"])
set_text(tag, "FORMAÇÃO DE GRUPOS")
pill.width = Emu(int(pill.width * 18 / 31) + 60000)
tag.width = Emu(int(tag.width * 18 / 31) + 60000)
title = clone(shapes["Text 2"])
set_text(title, "Modelo de Tuckman: as Etapas de um Grupo")

iw, ih = Image.open(IMG).size
w = Emu(10868745)
h = Emu(int(w * ih / iw))
s.shapes.add_picture(IMG, Emu(661475), Emu(1350000), w, h)

s.shapes._spTree.append(copy.deepcopy(cap_src._element))
cap = s.shapes[-1]
paras = cap.text_frame._txBody.findall(A + "p")
runs = paras[0].findall(A + "r")
for r in runs[1:]:
    paras[0].remove(r)
runs[0].find(A + "t").text = (
    "Bruce Tuckman (1965) descreveu as etapas por que passa um grupo pequeno "
    "orientado para uma tarefa. Em 1977, com Mary Ann Jensen, acrescentou a dissolução.")
runs[0].find(A + "rPr").set("b", "0")
cap.top = Emu(1350000 + h + 150000)
cap.text_frame.auto_size = MSO_AUTO_SIZE.NONE

ns = s.notes_slide
if ns.notes_placeholder is None:
    ns.shapes._spTree.append(copy.deepcopy(ref.notes_slide.notes_placeholder._element))
ns.notes_text_frame.text = (
    "• Um grupo de trabalho não é uma coleção de indivíduos ao acaso: os membros interagem, "
    "influenciam-se e desenvolvem uma forma própria de trabalhar para alcançar metas partilhadas\n\n"
    "• Formação: o grupo forma-se; cada pessoa identifica as metas partilhadas e sente que pertence "
    "ao grupo; os conflitos são evitados\n\n"
    "• Confrontação: é a etapa com mais conflitos; definem-se os papéis e as responsabilidades de "
    "cada um e surge uma liderança\n\n"
    "• Normatização: o grupo define o processo de trabalho; há menos conflito, porque os membros já "
    "conhecem as capacidades uns dos outros; o líder tem um papel forte\n\n"
    "• Atuação: produtividade alta e pouco conflito; o grupo trabalha por conta própria e ajusta o "
    "processo quando é preciso; a liderança intervém menos\n\n"
    "• Dissolução: o grupo desfaz-se porque a tarefa terminou ou porque desistiu; reconhece-se o que "
    "foi feito e os membros separam-se\n\n"
    "• Os ícones mostram o movimento dos membros: convergem, chocam, alinham-se, avançam juntos e, "
    "no fim, separam-se\n\n"
    "• Pergunta para a turma: em que etapa está o vosso grupo de trabalho desta unidade curricular?")

lst = p.slides._sldIdLst
new = list(lst)[-1]
lst.remove(new)
lst.insert(idx + 1, new)
p.save(DECK)
print("inserido como slide", idx + 2, "de", len(p.slides))

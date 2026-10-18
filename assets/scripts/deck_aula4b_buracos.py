import sys
from pptx import Presentation

# Aula 4b: buraco estrutural = ligação que falta entre dois contactos de uma
# pessoa (Burt, 1992); E deixa de ser "o único corretor" (C também atravessa
# buracos), passa a "o corretor entre os dois grupos".
DECK = sys.argv[1]

p = Presentation(DECK)
done = 0
for s in p.slides:
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            runs = para.runs
            texts = [r.text for r in runs]
            if texts[:1] == ["E é o único corretor"]:
                runs[0].text = "E é o corretor entre os dois grupos"
                runs[1].text = ": as ligações C–E e E–F são as únicas pontes entre eles"
                done += 1
            if texts[:1] == ["Buracos estruturais (Burt, 1992)"]:
                runs[1].text = (" — uma ligação que falta entre dois contactos de uma pessoa: "
                                "ambos a conhecem, mas não se conhecem entre si. Quem fica no meio, o ")
                runs[2].text = "corretor"
                runs[3].text = (", acede primeiro à informação dos dois lados; é o que a intermediação "
                                "mede. Na turma, C e F não se conhecem e E está no meio: se E faltar, "
                                "os dois grupos perdem toda a ligação.")
                done += 1
    if s.has_notes_slide:
        for para in s.notes_slide.notes_text_frame.paragraphs:
            for r in para.runs:
                old = ("• Burt (1992): quem ocupa o buraco estrutural entre dois grupos recebe primeiro "
                       "a informação dos dois lados; é a intermediação. Na turma, E")
                if old in r.text:
                    r.text = r.text.replace(old, (
                        "• Burt (1992): um buraco estrutural é uma ligação que falta entre dois contactos "
                        "de uma pessoa (ambos a conhecem, mas não se conhecem). Não é qualquer ligação em "
                        "falta: tem de haver alguém no meio. Quem está no meio é o corretor e recebe "
                        "primeiro a informação dos dois lados; é a intermediação\n\n"
                        "• Na turma: E tem dois contactos, C e F, que não se conhecem; E atravessa esse "
                        "buraco. C também atravessa buracos (é o único do seu grupo ligado a E) e tem a "
                        "intermediação mais alta. Os dois estão no único caminho entre os grupos: se C ou E "
                        "faltarem, a rede parte-se. A diferença: C está bem integrado no seu grupo; E não tem "
                        "mais nada além do buraco, é o corretor no sentido mais puro\n\n"
                        "• Ligar ao agrupamento: o coeficiente conta os pares de contactos ligados; os "
                        "buracos são os pares que não estão. E tem coeficiente 0: o seu único par é um buraco"))
                    done += 1
assert done == 3, done
p.save(DECK)
print("ok")

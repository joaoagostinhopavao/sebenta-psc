import copy
import sys
from pptx import Presentation
from pptx.util import Emu

# Aula 5: o slide "Cooperação: A Atividade como Elemento Central" passa a
# ter as três características das tarefas em cartões, como os outros
# slides da ontologia; o texto perde os travessões e o "Se articulam".
# Cartão e estilos de título/corpo copiados do slide da Comunicação.
# Só com o deck fechado.
DECK = sys.argv[1]
CARTOES = [
    ("Recursos", "Cada tarefa requer recursos específicos para a sua execução."),
    ("Execução individual ou partilhada",
     "Uma tarefa pode ser feita por um participante ou por vários em simultâneo."),
    ("Articulação com o plano",
     "Cada tarefa articula-se com o plano de trabalho definido na coordenação."),
]
SVG = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

p = Presentation(DECK)


def slide_com(titulo):
    return next(s for s in p.slides
                if any(sh.has_text_frame and sh.text_frame.text == titulo for sh in s.shapes))


coop = slide_com("Cooperação: A Atividade como Elemento Central")
com = slide_com("Comunicação: A Mensagem como Elemento Central")
if any(sh.shape_type == 13 for sh in coop.shapes):
    sys.exit("o slide já tem cartões")
sc = {s.name: s for s in com.shapes}
sh = {s.name: s for s in coop.shapes}

# texto: sem travessões, e a lista passa para os cartões
t3 = sh["Text 3"].text_frame.paragraphs
t3[0].runs[2].text = (", realizada por um grupo com uma determinada configuração geográfica "
                      "(colocalizada ou remota) e um determinado acoplamento entre participantes.")
t3[1].runs[4].text = ", composto por artefactos concretos. A produção decompõe-se em tarefas:"
lista = sh["Text 4"]._element
ancora = lista.getprevious()
lista.getparent().remove(lista)

cartao = next(s for s in com.shapes if s.shape_type == 13)
svg_rid = cartao._element.find(".//" + SVG).get(R_EMBED)
rid = coop.part.relate_to(com.part.related_part(svg_rid), com.part.rels[svg_rid].reltype)

topo, passo = 2950000, 1000000
dy_titulo, dy_corpo = 168870, 538857
proximo_id = 950
for i, (titulo, corpo) in enumerate(CARTOES):
    y = topo + i * passo
    els = []
    pic = copy.deepcopy(cartao._element)
    pic.find(".//" + SVG).set(R_EMBED, rid)
    pic.find(".//" + A + "off").set("y", str(y))
    els.append(pic)
    for modelo, texto, dy in [(sc["Text 4"], titulo, dy_titulo), (sc["Text 5"], corpo, dy_corpo)]:
        el = copy.deepcopy(modelo._element)
        el.find(".//" + A + "off").set("y", str(y + dy))
        runs = el.findall(".//" + A + "r")
        runs[0].find(A + "t").text = texto
        for r in runs[1:]:
            r.getparent().remove(r)
        els.append(el)
    for el in els:
        el.find(".//" + P + "cNvPr").set("id", str(proximo_id))
        proximo_id += 1
        ancora.addnext(el)
        ancora = el

sh["Text 5"].top = Emu(topo + 2 * passo + cartao.height + 90000)
p.save(DECK)
print("ok", list(p.slides).index(coop) + 1)

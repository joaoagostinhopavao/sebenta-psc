import copy
import sys
from pptx import Presentation
from pptx.util import Emu

SRC, OUT, IMG = sys.argv[1], sys.argv[2], sys.argv[3]
p = Presentation(SRC)
s8 = p.slides[7]
shapes8 = {sh.name: sh for sh in s8.shapes}


def new_slide():
    s = p.slides.add_slide(s8.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def copy_shape(slide, name, top=None):
    el = copy.deepcopy(shapes8[name]._element)
    slide.shapes._spTree.append(el)
    sh = slide.shapes[-1]
    if top is not None:
        sh.top = Emu(top)
    return sh


def set_paragraphs(sh, texts, size=None, bold_last=False):
    txBody = sh.text_frame._txBody
    paras = txBody.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}p")
    model = paras[0]
    for para in paras:
        txBody.remove(para)
    for t in texts:
        para = copy.deepcopy(model)
        runs = para.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
        for r in runs[1:]:
            para.remove(r)
        runs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = t
        rPr = runs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
        if size:
            rPr.set("sz", str(size))
        rPr.set("b", "1" if (bold_last and t is texts[-1]) else "0")
        txBody.append(para)


def header(slide, title, dy):
    copy_shape(slide, "Shape 0", shapes8["Shape 0"].top - dy)
    copy_shape(slide, "Text 1", shapes8["Text 1"].top - dy)
    t = copy_shape(slide, "Text 2", shapes8["Text 2"].top - dy)
    set_paragraphs(t, [title])
    t.text_frame.paragraphs[0].runs[0].font.bold = None


def notes(slide, text):
    ns = slide.notes_slide
    if ns.notes_placeholder is None:
        src = s8.notes_slide.notes_placeholder._element
        ns.shapes._spTree.append(copy.deepcopy(src))
    ns.notes_text_frame.text = text


# Slide A: a situação
a = new_slide()
header(a, "Exemplo: a Viagem de Finalistas", 0)
body = copy_shape(a, "Text 3")
body.width = Emu(10868745)
set_paragraphs(body, [
    "Os 40 finalistas de uma turma querem fazer uma viagem de finalistas em maio.",
    "1. Quatro colegas, a Ana, o Bruno, a Carla e o Duarte, formam a comissão que vai organizar a viagem.",
    "2. A comissão cria um grupo de WhatsApp com toda a turma e uma folha de cálculo partilhada para registar quem já pagou.",
    "3. A turma combina três regras: cada decisão é votada numa sondagem do WhatsApp, a viagem não pode custar mais de 350 € por pessoa e a quota paga-se até 15 de março.",
    "4. A comissão reparte o trabalho: a Ana trata do alojamento, o Bruno do transporte, a Carla dos pagamentos e o Duarte das atividades.",
    "5. Em abril, a viagem está reservada e paga.",
    "Quem é o sujeito? Qual é o objeto? Quais são os artefactos, a comunidade, as regras, a divisão de trabalho e o resultado?",
], size=1600, bold_last=True)
notes(a, "• Ler a situação passo a passo com a turma\n\n"
         "• Antes de passar ao slide seguinte, pedir aos alunos que identifiquem os sete elementos do modelo: sujeito, objeto, artefactos, comunidade, regras, divisão de trabalho e resultado\n\n"
         "• No teste, a pergunta segue este formato: uma situação descrita em texto e a identificação dos elementos do modelo, com uma frase de justificação para cada um")

# Slide B: o modelo preenchido
b = new_slide()
dy = shapes8["Shape 0"].top - 462359
header(b, "Exemplo: o Modelo Preenchido", dy)
img_h = 4400000
from PIL import Image
iw, ih = Image.open(IMG).size
img_w = int(img_h * iw / ih)
img_top = 1450000
b.shapes.add_picture(IMG, Emu((p.slide_width - img_w) // 2), Emu(img_top), Emu(img_w), Emu(img_h))
line = copy_shape(b, "Text 3", img_top + img_h + 120000)
line.left = Emu(661475)
line.width = Emu(10868745)
line.height = Emu(500000)
set_paragraphs(line, [
    "Ação reversa do artefacto: com as sondagens, a comissão deixa de decidir em longas discussões e passa a decidir por votação.",
], size=1500)
notes(b, "• Sujeito: a comissão de 4 finalistas, é quem age; não é a turma inteira, que só vota e paga\n\n"
         "• Objeto: a organização da viagem, aquilo sobre que a comissão trabalha\n\n"
         "• Resultado: a viagem reservada e paga, aquilo em que o objeto se transforma no fim\n\n"
         "• Artefactos: o grupo de WhatsApp com as sondagens e a folha de cálculo partilhada; a comissão não age diretamente sobre a viagem, age através destas ferramentas\n\n"
         "• Comunidade: os 40 finalistas da turma, o grupo onde a comissão está inserida\n\n"
         "• Regras: decisões por votação, máximo de 350 € por pessoa, quota paga até 15 de março\n\n"
         "• Divisão de trabalho: Ana trata do alojamento, Bruno do transporte, Carla dos pagamentos, Duarte das atividades\n\n"
         "• Ação reversa: o artefacto também muda o sujeito; a sondagem do WhatsApp muda a forma como a comissão decide")

# mover os dois slides novos para depois do slide 8
lst = p.slides._sldIdLst
ids = list(lst)
new_ids = ids[-2:]
for el in new_ids:
    lst.remove(el)
for i, el in enumerate(new_ids):
    lst.insert(8 + i, el)

p.save(OUT)
print(len(p.slides))

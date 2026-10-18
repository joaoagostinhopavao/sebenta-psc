import copy
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

# Aula 5, slide "Formação de Grupos": cartão "Crença ≠ Confiança" na metade
# direita, com o mesmo aspeto dos cartões Confiança/Motivação.
DECK = sys.argv[1]
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

BODY = [
    ("Crença: ", "o que o participante pensa sobre um colega, por exemplo, que é competente."),
    ("Confiança: ", "a disposição para depender desse colega; assenta na crença."),
    ("Exemplo: ", "acreditar que um colega é competente e, mesmo assim, não lhe entregar "
                  "uma parte do trabalho. Há crença, mas não há confiança."),
]

p = Presentation(DECK)
slide = next(s for s in p.slides
             if any(sh.has_text_frame and sh.text_frame.text.startswith("Formação de Grupos: O Participante")
                    for sh in s.shapes))
shapes = {sh.name: sh for sh in slide.shapes}
assert not any(sh.has_text_frame and "Crença ≠ Confiança" in sh.text_frame.text for sh in slide.shapes)

L, T, W, H = 6846716, 1531243, 4689755, 2450000
card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(L), Emu(T), Emu(W), Emu(H))
card.adjustments[0] = 0.02
card.fill.solid()
card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
card.line.color.rgb = RGBColor(0xC5, 0xD2, 0xCF)
card.line.width = Pt(1.5)
card.shadow.inherit = False
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(L), Emu(T), Emu(76200), Emu(H))
bar.fill.solid()
bar.fill.fore_color.rgb = RGBColor(0x43, 0x70, 0x66)
bar.line.fill.background()
bar.shadow.inherit = False

title = copy.deepcopy(shapes["Text 4"]._element)
slide.shapes._spTree.append(title)
t = slide.shapes[-1]
t.left, t.top = Emu(L + 245064), Emu(T + 168870)
t.text_frame.paragraphs[0].runs[0].text = "Crença ≠ Confiança"

body_el = copy.deepcopy(shapes["Text 5"]._element)
slide.shapes._spTree.append(body_el)
b = slide.shapes[-1]
b.left, b.top = Emu(L + 245064), Emu(T + 538857)
b.width, b.height = Emu(W - 413931), Emu(H - 700000)
txBody = b.text_frame._txBody
bodyPr = txBody.find(A + "bodyPr")
for child in list(bodyPr):
    bodyPr.remove(child)
model = txBody.findall(A + "p")[0]
for para in txBody.findall(A + "p"):
    txBody.remove(para)
for k, (label, rest) in enumerate(BODY):
    para = copy.deepcopy(model)
    if k:
        pPr = para.find(A + "pPr")
        spc = pPr.makeelement(A + "spcBef", {})
        pts = spc.makeelement(A + "spcPts", {"val": "900"})
        spc.append(pts)
        pPr.insert(1, spc)
    r0 = para.findall(A + "r")[0]
    r1 = copy.deepcopy(r0)
    r0.find(A + "t").text = label
    r0.find(A + "rPr").set("b", "1")
    r1.find(A + "t").text = rest
    r1.find(A + "rPr").set("b", "0")
    r0.addnext(r1)
    txBody.append(para)

p.save(DECK)
print("ok")

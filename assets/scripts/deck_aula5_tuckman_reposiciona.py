import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

# Aula 5: o slide do Modelo de Tuckman sai do meio das ontologias e passa
# para logo a seguir ao "3C vs. Comando e Controlo"; imagem com a linha 3C
# e legenda sobre a mudança do grupo ao longo do tempo.
DECK = sys.argv[1]
IMG = "/Users/jpavao/QProjects/Sebenta-PSC/assets/images/cap04/slides/tuckman-slide.png"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

CAPTION = ("Um grupo não é estático: muda ao longo do tempo, da formação à dissolução. "
           "Em cada etapa, cada um dos três C tem uma importância diferente: a comunicação "
           "pesa mais na formação e na confrontação, a coordenação na normatização e a "
           "cooperação na atuação.")

NOTES = (
    "• Ideia principal: um grupo não é estático, muda ao longo do tempo; o Modelo 3C diz o que a "
    "colaboração exige, o Modelo de Tuckman diz como o grupo evolui\n\n"
    "• Bruce Tuckman (1965) descreveu quatro etapas por que passa um grupo pequeno orientado para "
    "uma tarefa; em 1977, com Mary Ann Jensen, acrescentou a dissolução\n\n"
    "• Formação: o grupo forma-se; cada pessoa identifica as metas partilhadas e sente que pertence "
    "ao grupo; os conflitos são evitados. Pesa a comunicação: conhecer os colegas e partilhar metas\n\n"
    "• Confrontação: é a etapa com mais conflitos; definem-se os papéis e as responsabilidades de "
    "cada um e surge uma liderança. Pesa a comunicação: negociar papéis\n\n"
    "• Normatização: o grupo define o processo de trabalho; há menos conflito e o líder tem um papel "
    "forte. Pesa a coordenação: processo, prazos, distribuição das tarefas\n\n"
    "• Atuação: produtividade alta e pouco conflito; o grupo trabalha por conta própria e a "
    "liderança intervém menos. Pesa a cooperação: produzir juntos no espaço partilhado\n\n"
    "• Dissolução: o grupo desfaz-se porque a tarefa terminou ou porque desistiu; reconhece-se o "
    "que foi feito e os membros separam-se\n\n"
    "• Os três C estão presentes em todas as etapas; muda qual deles exige mais atenção e, por "
    "isso, as ferramentas de que o grupo mais precisa. Esta ligação aos 3C combina os dois "
    "modelos, não é do Tuckman original\n\n"
    "• Exemplo: a comissão da viagem de finalistas oferece-se e combina metas (formação), discute "
    "quem faz o quê (confrontação), acorda votar por sondagem e registar pagamentos na folha "
    "(normatização), trabalha sem precisar de ser lembrada (atuação) e desfaz-se depois da viagem "
    "(dissolução)")


def title_of(slide):
    for sh in slide.shapes:
        if sh.name == "Text 2" and sh.has_text_frame:
            return sh.text_frame.text
    return ""


def set_first_run(sh, text):
    paras = sh.text_frame._txBody.findall(A + "p")
    for extra in paras[1:]:
        extra.getparent().remove(extra)
    runs = paras[0].findall(A + "r")
    for r in runs[1:]:
        paras[0].remove(r)
    runs[0].find(A + "t").text = text


p = Presentation(DECK)
slides = list(p.slides)
tk = next(s for s in slides if title_of(s).startswith("Modelo de Tuckman"))
c2 = next(s for s in slides if title_of(s).startswith("3C vs. Comando e Controlo"))

for sh in list(tk.shapes):
    if sh.name == "Text 1":
        set_first_run(sh, "MODELO DE TUCKMAN")
    elif sh.shape_type == 13:
        sh._element.getparent().remove(sh._element)

iw, ih = Image.open(IMG).size
top, h = 1300000, 4550000
w = int(h * iw / ih)
tk.shapes.add_picture(IMG, Emu((p.slide_width - w) // 2), Emu(top), Emu(w), Emu(h))

cap = next(sh for sh in tk.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Bruce Tuckman"))
set_first_run(cap, CAPTION)
cap.top = Emu(top + h + 100000)
cap.height = Emu(650000)

tk.notes_slide.notes_text_frame.text = NOTES

lst = p.slides._sldIdLst
ids = list(lst)
el = ids[slides.index(tk)]
lst.remove(el)
lst.insert(list(lst).index(ids[slides.index(c2)]) + 1, el)

p.save(DECK)
for i, s in enumerate(p.slides, 1):
    if 12 <= i <= 19:
        print(i, title_of(s))

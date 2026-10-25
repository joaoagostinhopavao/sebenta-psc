import sys
from pptx import Presentation

# Aula 5: notas do orador revistas na forma (PT-PT, sem travessões, sem
# anglicismos desnecessários, instruções no infinitivo). O conteúdo mantém-se.
# As chaves são o início do título de cada slide, para não depender da ordem.
# Só com o deck fechado.
DECK = sys.argv[1]

NOTAS = {
    "Teorias e modelos formais": (
        "• Explorar os fundamentos teóricos que sustentam a colaboração\n\n"
        "• Percorrer diferentes modelos, da Teoria dos Jogos até modelos mais complexos, como a ontologia\n\n"
        "• Estas teorias ajudam a compreender como as pessoas trabalham juntas e como desenhar sistemas que facilitem essa colaboração\n\n"
        "• Ver como aplicar estes conceitos em plataformas colaborativas reais\n\n"
        "• Corresponde ao Capítulo 4 da sebenta"
    ),
    "Contexto e Objetivos": (
        "• Mudança de foco: deixam-se as plataformas sociais (redes, métricas, folksonomia) e entra-se no bloco das plataformas colaborativas\n\n"
        "• A diferença é importante: já não se estuda a rede em si, mas o trabalho em grupo que ela pode suportar\n\n"
        "• Antes de desenhar ou avaliar qualquer sistema colaborativo, é preciso perceber que a colaboração pode ser modelada: existem teorias e modelos formais que explicam porque é que as pessoas colaboram\n\n"
        "• Quatro pilares:\n\n"
        "Teoria dos Jogos: porque é racional colaborar\n\n"
        "Teoria da Atividade: como se realiza uma atividade mediada\n\n"
        "Modelo 3C: as três dimensões essenciais, comunicação, coordenação e cooperação\n\n"
        "Ontologia de Colaboração: uma base terminológica comum para falar de colaboração\n\n"
        "• Estes modelos ajudam a perceber onde a colaboração tende a falhar e como desenhar sistemas que funcionem melhor"
    ),
    "Porque Modelar": (
        "• Um modelo científico é uma representação abstrata que ajuda a explicar, analisar e prever fenómenos de forma sistemática; pode ser uma fórmula, um diagrama ou um esquema conceptual\n\n"
        "• Modelar a colaboração é ir além da intuição:\n\n"
        "Em vez de decidir por instinto como apoiar um grupo, analisa-se que tipo de trabalho está realmente a ser feito\n\n"
        "Usam-se os mesmos conceitos para comparar situações diferentes\n\n"
        "As escolhas de sistemas colaborativos passam a assentar numa análise rigorosa\n\n"
        "• Os quatro benefícios principais:\n\n"
        "Explicar: perceber porque é que as pessoas colaboram, ou porque não colaboram\n\n"
        "Analisar: comparar contextos distintos com uma linguagem comum\n\n"
        "Prever: identificar onde a colaboração pode falhar antes de falhar\n\n"
        "Desenhar: construir ou escolher sistemas colaborativos com fundamento"
    ),
    "Decisão Estratégica e Colaboração": (
        "• Introdução à Teoria dos Jogos: o estudo de situações em que o resultado de cada pessoa depende não só da sua decisão, mas também das decisões dos outros\n\n"
        "• Contexto histórico: raízes no século XVIII, mas a teoria consolidou-se em 1944 com von Neumann e Morgenstern e, depois, com o conceito de equilíbrio de Nash, nos anos 1950\n\n"
        "• Jogo de soma nula:\n\n"
        "O ganho de um é exatamente a perda do outro\n\n"
        "A soma total é sempre zero\n\n"
        "Exemplos: xadrez e póquer, jogos puramente competitivos\n\n"
        "• Jogo de soma não nula:\n\n"
        "O resultado global pode variar\n\n"
        "Todos podem ganhar juntos ou todos podem perder juntos\n\n"
        "Exemplo importante: o Dilema dos Prisioneiros, que mostra como a colaboração, ou a falta dela, afeta o resultado coletivo\n\n"
        "• Transição: estes dois tipos de jogos representam cenários muito diferentes e exigem estratégias distintas"
    ),
    "O Dilema dos Prisioneiros": (
        "• Apresentar o Dilema dos Prisioneiros, um clássico da Teoria dos Jogos\n\n"
        "• Contexto: dois cúmplices de um crime, detidos e interrogados em separado, sem comunicação entre si\n\n"
        "• Explicar os três cenários possíveis:\n\n"
        "Ambos em silêncio → 6 meses cada\n\n"
        "Um confessa e o outro cala-se → quem confessa sai em liberdade, o outro cumpre 10 anos\n\n"
        "Ambos confessam → 6 anos cada\n\n"
        "• Destacar o paradoxo central: cada prisioneiro, pensando racionalmente no seu próprio interesse, escolhe confessar. Mas, quando ambos o fazem, o resultado é pior do que se tivessem colaborado e ficado em silêncio\n\n"
        "• Reforçar a lição: a ação racional individual não leva ao melhor resultado coletivo. Isto mostra porque é que a cooperação é difícil mesmo quando seria benéfica para ambos"
    ),
    "Tit-for-Tat": (
        "• Introduzir o contexto: quando o Dilema dos Prisioneiros se repete entre os mesmos jogadores, tudo muda\n\n"
        "• Referir Robert Axelrod e as suas competições de estratégias nos anos 70\n\n"
        "• Destacar que o tit-for-tat foi a estratégia vencedora, proposta por Anatol Rapoport\n\n"
        "• Explicar as três regras simples:\n\n"
        "Coopera primeiro: nunca sejas o primeiro a trair\n\n"
        "Retalia se traído: responde à traição na ronda seguinte\n\n"
        "Perdoa e volta a cooperar: depois de retaliar, está disposto a retomar a colaboração\n\n"
        "• O ponto-chave: saber que se vai voltar a interagir com o mesmo adversário torna a traição muito custosa\n\n"
        "• Conclusão: o tit-for-tat revela-se mais eficaz a longo prazo do que as estratégias puramente egoístas\n\n"
        "• Implicação importante: isto explica como a colaboração pode surgir naturalmente, sem ser preciso uma autoridade central que a imponha"
    ),
    "A Tragédia dos Comuns": (
        "• Apresentar o conceito de Tragédia dos Comuns, descrito por Garrett Hardin em 1968\n\n"
        "• Explicar o dilema: um recurso comum e limitado, partilhado por vários participantes\n\n"
        "• Usar o exemplo do pasto:\n\n"
        "Cada pastor lucra individualmente ao acrescentar mais animais\n\n"
        "O custo do desgaste é dividido por todos\n\n"
        "Quando todos pensam assim, o recurso é destruído e todos perdem\n\n"
        "• Destacar os três elementos-chave:\n\n"
        "Ganho privado: o benefício fica inteiramente com quem decide\n\n"
        "Custo partilhado: o custo é distribuído pelo grupo todo\n\n"
        "Solução: regras explícitas e penalizações são essenciais\n\n"
        "• Reforçar que, sem mecanismos de controlo e normas, o interesse individual leva à ruína coletiva"
    ),
    "Como se Realiza uma Atividade Mediada?": (
        "• A Teoria da Atividade explica como se realizam atividades no dia a dia, sozinho ou em grupo\n\n"
        "A unidade básica é a atividade: um sujeito age sobre um objeto para atingir um objetivo\n\n"
        "• O elemento-chave: a mediação por artefactos\n\n"
        "Podem ser físicos (computador, lápis) ou cognitivos (linguagem, conceitos)\n\n"
        "O artefacto não atua apenas sobre o objeto: também transforma a forma como o sujeito pensa e resolve problemas\n\n"
        "• Na dimensão coletiva:\n\n"
        "O sujeito faz parte de uma comunidade\n\n"
        "A comunidade tem regras que orientam as ações\n\n"
        "Existe uma divisão de trabalho que define responsabilidades\n\n"
        "• Referir o diagrama: mostra como sujeito, objeto, artefactos, comunidade, regras e divisão de trabalho se relacionam num sistema integrado"
    ),
    "3C vs. Comando e Controlo": (
        "• Recordar as três dimensões essenciais para uma colaboração real\n\n"
        "• Comunicação: negociação entre participantes, troca de ideias\n\n"
        "• Coordenação: organização de tarefas distribuídas, sincronização de esforços\n\n"
        "• Cooperação: trabalho conjunto num espaço partilhado, objetivos comuns\n\n"
        "• Contraste com Comando e Controlo (C2): modelo centralizado, típico das estruturas militares e da produção industrial clássica\n\n"
        "• Diferença-chave: no C2, uma autoridade central decide tudo e supervisiona diretamente, o que elimina a necessidade de comunicação e cooperação entre quem executa\n\n"
        "• Mensagem central: para haver verdadeira colaboração, e não apenas execução de ordens, as três dimensões do 3C têm de estar presentes ao mesmo tempo\n\n"
        "• Implicação: as organizações que querem colaboração real precisam de abandonar a lógica do C2 e de criar espaço para a negociação, a coordenação distribuída e a cooperação genuína"
    ),
    "O que é uma Ontologia": (
        "• Começar por explicar que uma ontologia é a descrição formal de um domínio, negociada pela comunidade para um objetivo específico\n\n"
        "• Destacar o ponto-chave: cria uma base terminológica comum, que melhora a comunicação entre os profissionais da área e permite que os sistemas computacionais compreendam e tratem a informação sobre esse domínio\n\n"
        "• Apresentar a Ontologia de Colaboração, baseada no modelo de Ellis e colaboradores\n\n"
        "• Explicar que se organiza em quatro partes principais:\n\n"
        "Formação de Grupos (conceito central: Participante): como as pessoas se reúnem\n\n"
        "Comunicação (conceito central: Mensagem): como trocam informação\n\n"
        "Coordenação (conceito central: Plano de Trabalho): como organizam o trabalho\n\n"
        "Cooperação (conceito central: Atividade): como executam as tarefas em conjunto\n\n"
        "• Reforçar que estes quatro pilares funcionam em conjunto para estruturar a colaboração de forma clara e sistemática"
    ),
    "Formação de Grupos: O Participante": (
        "• O participante é o elemento central na formação de grupos: assume tarefas, integra-se no grupo e participa na negociação dos objetivos comuns\n\n"
        "• Duas condições essenciais para que isso funcione:\n\n"
        "Confiança: a disposição para depender dos outros, aceitando ficar exposto ao que fizerem, com base na expectativa de que cumprirão os compromissos assumidos (Mayer, Davis e Schoorman, 1995)\n\n"
        "Não confundir com as crenças, outro atributo do participante no diagrama: as crenças são o que o participante pensa (por exemplo, que um colega é competente); a confiança é a disposição para depender dele e assenta nessas crenças. Pode haver crença sem confiança: acreditar que o colega é competente e, mesmo assim, não lhe entregar parte do trabalho\n\n"
        "Motivação: sem ela, o participante não se compromete nem contribui ativamente\n\n"
        "• O grupo realiza uma atividade que gera um produto, e esse produto é também um objetivo partilhado"
    ),
    "Coordenação: O Plano de Trabalho": (
        "• O plano de trabalho é o centro da coordenação: resulta da negociação e é suportado pelo sistema colaborativo\n\n"
        "• Três elementos principais definem o plano:\n\n"
        "Papéis e Acoplamento: deixar claro quem faz o quê e como as atividades se ligam\n\n"
        "Políticas e Prazos: estabelecer as regras do jogo e quando deve acontecer cada entrega\n\n"
        "Alocação de Recursos: definir o que cada tarefa precisa e distribuir os recursos de forma adequada\n\n"
        "• O acompanhamento fecha o ciclo: verifica se os prazos estão a ser cumpridos e em que estado está cada atividade\n\n"
        "• Sem um plano claro e bem acompanhado, a coordenação desfaz-se"
    ),
    "Cooperação: A Atividade": (
        "• A cooperação tem a atividade no centro: é o que o grupo faz em conjunto\n\n"
        "Pode ser colocalizada ou remota\n\n"
        "Depende de como os participantes estão acoplados\n\n"
        "• A atividade decorre num espaço de trabalho partilhado e produz algo concreto: um produto feito de artefactos reais\n\n"
        "• O trabalho divide-se em tarefas que:\n\n"
        "Precisam de recursos específicos\n\n"
        "Podem ser individuais ou feitas por vários em simultâneo\n\n"
        "Têm de estar alinhadas com o plano de coordenação\n\n"
        "• É aqui que a cooperação se torna produção real: deixa de ser apenas intenção e passa a ser algo tangível"
    ),
    "A Ontologia de Colaboração Completa": (
        "• A ontologia de colaboração completa sintetiza quatro elementos essenciais\n\n"
        "• Participantes motivados: têm de estar empenhados e comprometidos com o objetivo comum\n\n"
        "• Mensagens: a comunicação é fundamental para negociar e assumir compromissos entre os participantes\n\n"
        "• Plano: organiza as tarefas e coordena o trabalho de forma estruturada\n\n"
        "• Artefactos: os produtos concretos criados no espaço partilhado\n\n"
        "• Reforçar: se faltar qualquer um destes quatro elementos, não há verdadeira colaboração\n\n"
        "• Para quem desenha sistemas colaborativos, esta ontologia funciona como lista de verificação: garante que aspetos como a confiança entre participantes e a existência de um plano explícito não são esquecidos no desenho"
    ),
    "Síntese": (
        "• Teoria dos Jogos: começa-se por perceber como as pessoas decidem em situações de conflito. O que é racional para cada um pode ser mau para todos (Dilema dos Prisioneiros), mas, em jogos repetidos, a colaboração consegue surgir sem ninguém a impor\n\n"
        "• Teoria da Atividade: muda o foco do indivíduo para o sistema todo. Cada ação passa por ferramentas, regras da comunidade e divisão de tarefas; tudo está interligado\n\n"
        "• Modelo 3C: decompõe a colaboração em três dimensões que se alimentam mutuamente: comunicação (negociar o que fazer), coordenação (organizar como fazer) e cooperação (fazer em conjunto no mesmo espaço)\n\n"
        "• Ontologia de Colaboração: junta tudo isto numa linguagem comum, com quatro peças-chave: quem participa, o que se comunica, como se planeia o trabalho e que atividades se desenvolvem"
    ),
    "Para Refletir": (
        "• Convidar a turma a refletir sobre um trabalho de grupo recente\n\n"
        "Pedir que identifiquem os três C: comunicação, coordenação e cooperação\n\n"
        "Qual foi o mais fraco? Que problemas concretos causou?\n\n"
        "• Segunda reflexão, sobre relações duradouras\n\n"
        "Podem ser profissionais, de amizade ou entre organizações\n\n"
        "A colaboração mantém-se porque se espera voltar a interagir\n\n"
        "Reconhecem a estratégia do tit-for-tat (reciprocidade)?\n\n"
        "• Terceira questão, sobre a Tragédia dos Comuns\n\n"
        "Pedir um exemplo concreto que conheçam\n\n"
        "Que regras ou penalizações existem para a evitar?\n\n"
        "Formais ou informais, ambas contam\n\n"
        "• Dar tempo para reflexão pessoal: estas questões servem para interiorizar os conceitos e não precisam de resposta em voz alta"
    ),
}

p = Presentation(DECK)
usadas = set()
for n, s in enumerate(p.slides, 1):
    textos = [sh.text_frame.text for sh in s.shapes
              if sh.has_text_frame and sh.name != "NumeroSlide"]
    chave = next((k for k in NOTAS if any(t.startswith(k) for t in textos)), None)
    if chave:
        s.notes_slide.notes_text_frame.text = NOTAS[chave]
        usadas.add(chave)
        print(n, chave)
falta = set(NOTAS) - usadas
if falta:
    sys.exit(f"chaves sem slide: {falta}")
p.save(DECK)
